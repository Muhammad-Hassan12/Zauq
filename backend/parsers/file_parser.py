import io
import os
import re
import csv
import base64
import mimetypes
from typing import Dict, Any, List, Tuple, Optional
from pypdf import PdfReader
from docx import Document
import openpyxl

TEXT_EXTENSIONS = {
    ".txt", ".py", ".js", ".ts", ".html", ".css", ".json", ".yaml", ".yml",
    ".md", ".log", ".sh", ".sql", ".c", ".cpp", ".h", ".java", ".go", ".rs",
    ".php", ".xml", ".ini", ".env", ".toml", ".bat", ".ps1"
}

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}

AUDIO_EXTENSIONS = {".ogg", ".mp3", ".wav", ".m4a", ".aac", ".flac", ".opus", ".webm", ".oga"}

def parse_pdf(content_bytes: bytes) -> str:
    """Extracts text page-by-page from PDF bytes."""
    try:
        reader = PdfReader(io.BytesIO(content_bytes))
        pages_text = []
        total_len = 0
        for idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                pages_text.append(f"--- Page {idx + 1} ---\n{text.strip()}")
                total_len += len(text)
            if total_len > 120000:
                pages_text.append("\n... [PDF Content Truncated for Size Limits]")
                break
        return "\n\n".join(pages_text) if pages_text else "[PDF file contains no extractable text]"
    except Exception as e:
        return f"[PDF Parsing Error: {str(e)}]"

def parse_docx(content_bytes: bytes) -> str:
    """Extracts text and tables from Word (.docx) bytes."""
    try:
        doc = Document(io.BytesIO(content_bytes))
        output = []
        total_len = 0
        for p in doc.paragraphs:
            if p.text.strip():
                output.append(p.text.strip())
                total_len += len(p.text)
            if total_len > 120000:
                break

        if total_len <= 120000:
            for table_idx, table in enumerate(doc.tables):
                output.append(f"\n--- Table {table_idx + 1} ---")
                for row in table.rows:
                    row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    output.append("| " + " | ".join(row_cells) + " |")
                    total_len += len(" | ".join(row_cells))
                if total_len > 120000:
                    break

        if total_len > 120000:
             output.append("\n... [Docx Content Truncated for Size Limits]")

        return "\n".join(output) if output else "[Docx file contains no text]"
    except Exception as e:
        return f"[Docx Parsing Error: {str(e)}]"

def parse_excel(content_bytes: bytes, filename: str) -> str:
    """Extracts sheet data from Excel (.xlsx) or CSV bytes into Markdown tables."""
    ext = os.path.splitext(filename.lower())[1]
    if ext == ".csv":
        try:
            text = content_bytes.decode("utf-8", errors="replace")
            reader = csv.reader(io.StringIO(text))
            rows = list(reader)
            if not rows:
                return "[CSV file is empty]"
            markdown_rows = []
            markdown_rows.append("| " + " | ".join(rows[0]) + " |")
            markdown_rows.append("| " + " | ".join(["---"] * len(rows[0])) + " |")
            for r in rows[1:100]:  # Limit to first 100 rows
                markdown_rows.append("| " + " | ".join(r) + " |")
            return "\n".join(markdown_rows)
        except Exception as e:
            return f"[CSV Parsing Error: {str(e)}]"

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=True)
        sheets_output = []
        for sheet_name in wb.sheetnames[:3]:  # Max 3 sheets
            sheet = wb[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            sheets_output.append(f"--- Sheet: {sheet_name} ---")
            header = [str(c or "") for c in rows[0]]
            sheets_output.append("| " + " | ".join(header) + " |")
            sheets_output.append("| " + " | ".join(["---"] * len(header)) + " |")
            for r in rows[1:50]:  # Limit to 50 rows per sheet
                row_str = [str(c or "") for c in r]
                sheets_output.append("| " + " | ".join(row_str) + " |")
        return "\n\n".join(sheets_output) if sheets_output else "[Excel workbook is empty]"
    except Exception as e:
        return f"[Excel Parsing Error: {str(e)}]"

def parse_text(content_bytes: bytes) -> str:
    """Decodes plain text, code, JSON, and log files."""
    try:
        return content_bytes.decode("utf-8", errors="replace")
    except Exception as e:
        return f"[Text Decoding Error: {str(e)}]"

def parse_attachment(content_bytes: bytes, filename: str, content_type: str = "") -> Dict[str, Any]:
    """
    Main file parsing dispatcher. Detects file type (Images, Audio/Voice Notes, Documents, Text)
    and returns parsed content payload.
    """
    ext = os.path.splitext(filename.lower())[1]

    # 1. Multimodal Images
    if ext in IMAGE_EXTENSIONS or content_type.startswith("image/"):
        mime = content_type or f"image/{ext.replace('.', '')}"
        b64_str = base64.b64encode(content_bytes).decode("utf-8")
        return {
            "type": "image",
            "filename": filename,
            "mime_type": mime,
            "bytes_b64": b64_str,
            "content": f"[Attached Image: {filename}]"
        }

    # 2. Multimodal Audio / Discord Voice Notes
    if ext in AUDIO_EXTENSIONS or content_type.startswith("audio/"):
        mime = content_type or (f"audio/{ext.replace('.', '')}" if ext != ".oga" else "audio/ogg")
        if mime == "audio/opus" or ext == ".opus":
            mime = "audio/ogg"
        b64_str = base64.b64encode(content_bytes).decode("utf-8")
        return {
            "type": "audio",
            "filename": filename,
            "mime_type": mime,
            "bytes_b64": b64_str,
            "content": f"[Attached Voice Note / Audio Recording: {filename}]"
        }

    # 3. Documents & Spreadsheets
    if ext == ".pdf":
        extracted = parse_pdf(content_bytes)
    elif ext == ".docx":
        extracted = parse_docx(content_bytes)
    elif ext in [".xlsx", ".xls", ".csv"]:
        extracted = parse_excel(content_bytes, filename)
    elif ext in TEXT_EXTENSIONS or content_type.startswith("text/") or "json" in content_type:
        extracted = parse_text(content_bytes)
    else:
        # Fallback attempt to read as text
        extracted = parse_text(content_bytes)

    # Truncate extremely long document extractions to 120,000 chars (~30,000 tokens)
    if len(extracted) > 120000:
        extracted = extracted[:119900] + "\n... [Document Content Truncated for Prompt Length]"

    return {
        "type": "text",
        "filename": filename,
        "content": f"```\n{extracted}\n```"
    }

def extract_generated_files(response_text: str) -> Tuple[str, List[Dict[str, Any]]]:
    """
    Parses generated files from model output enclosed in:
    1. <zauq_file filename="app.py">...code...</zauq_file>
    2. Explicit fenced blocks with filename annotations like ```python:app.py ... ```
    
    Returns a tuple of (cleaned_response_text, list_of_file_dicts)
    where file dict contains {"filename": str, "content": str, "content_type": str, "bytes_b64": str}
    """
    if not response_text:
        return response_text, []

    files = []
    seen_filenames = set()

    # Helper to determine syntax highlighting language from filename
    def get_lang(fname: str) -> str:
        ext = os.path.splitext(fname.lower())[1]
        ext_map = {
            ".html": "html", ".htm": "html", ".py": "python", ".js": "javascript",
            ".ts": "typescript", ".json": "json", ".css": "css", ".sql": "sql",
            ".sh": "bash", ".md": "markdown", ".cpp": "cpp", ".c": "c",
            ".rs": "rust", ".go": "go", ".java": "java", ".yaml": "yaml", ".yml": "yaml"
        }
        return ext_map.get(ext, "")

    # Pattern 1: <zauq_file filename="...">...</zauq_file> or unclosed <zauq_file filename="...">...
    tag_pattern = r'<zauq_file\s+filename=["\']([^"\']+)["\']>(.*?)(?:</zauq_file>|\Z)'
    
    def tag_replacer(match: re.Match) -> str:
        filename = match.group(1).strip()
        code_content = match.group(2).strip()
        # If code_content is wrapped inside a fenced code block, strip the outer fences
        if code_content.startswith("```") and code_content.endswith("```"):
            lines = code_content.split("\n")
            if len(lines) >= 2:
                code_content = "\n".join(lines[1:-1]).strip()

        if filename and code_content and filename not in seen_filenames:
            seen_filenames.add(filename)
            mime, _ = mimetypes.guess_type(filename)
            b64_data = base64.b64encode(code_content.encode("utf-8")).decode("utf-8")
            files.append({
                "filename": filename,
                "content": code_content,
                "content_type": mime or "text/plain",
                "bytes_b64": b64_data
            })
            lang = get_lang(filename)
            
            # Show full code in chat if under 1200 chars, or first 25 lines preview if large
            if len(code_content) <= 1200:
                return f"\n\n📄 **Generated File:** `{filename}` *(downloadable file attached below)*\n```{lang}\n{code_content}\n```\n"
            else:
                preview_lines = code_content.split("\n")[:25]
                preview_snippet = "\n".join(preview_lines)
                return f"\n\n📄 **Generated File:** `{filename}` *(full file attached below)*\n```{lang}\n{preview_snippet}\n... [Full code in attached {filename}]\n```\n"
        elif filename and not code_content:
            return ""
        return ""

    cleaned_text = re.sub(tag_pattern, tag_replacer, response_text, flags=re.DOTALL)

    # Pattern 2: ```lang:filename.ext ... ``` if not already captured
    fenced_file_pattern = r'```([a-zA-Z0-9_\-\+]+):([a-zA-Z0-9_\-\.\/]+)\n(.*?)```'
    def fenced_replacer(match: re.Match) -> str:
        lang = match.group(1).strip()
        filename = match.group(2).strip()
        code_content = match.group(3).strip()
        if filename and code_content and filename not in seen_filenames:
            seen_filenames.add(filename)
            mime, _ = mimetypes.guess_type(filename)
            b64_data = base64.b64encode(code_content.encode("utf-8")).decode("utf-8")
            files.append({
                "filename": filename,
                "content": code_content,
                "content_type": mime or "text/plain",
                "bytes_b64": b64_data
            })
            if len(code_content) <= 1200:
                return f"\n\n📄 **Generated File:** `{filename}` *(downloadable file attached below)*\n```{lang}\n{code_content}\n```\n"
            else:
                preview_lines = code_content.split("\n")[:25]
                preview_snippet = "\n".join(preview_lines)
                return f"\n\n📄 **Generated File:** `{filename}` *(full file attached below)*\n```{lang}\n{preview_snippet}\n... [Full code in attached {filename}]\n```\n"
        return match.group(0)

    cleaned_text = re.sub(fenced_file_pattern, fenced_replacer, cleaned_text, flags=re.DOTALL)

    return cleaned_text.strip(), files
